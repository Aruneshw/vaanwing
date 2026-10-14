"""Extractors: each turns one source into located Segments (page / slide / timestamp / figure)."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from ..llm import LLMUnavailable, complete, enabled, image_block
from ...config import settings
from .chunker import Segment

log = logging.getLogger(__name__)


def kind_from_name(name: str) -> str:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in {"pdf"}:
        return "pdf"
    if ext in {"ppt", "pptx"}:
        return "slides"
    if ext in {"png", "jpg", "jpeg", "gif", "webp"}:
        return "image"
    if ext in {"mp4", "mov", "webm", "mkv"}:
        return "video"
    if ext in {"mp3", "wav", "m4a", "ogg"}:
        return "audio"
    if ext in {"docx"}:
        return "docx"
    return "notes"


def page_topic(text: str) -> str:
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for i, l in enumerate(lines):
        if re.match(r"^TOPIC \d+", l) and i + 1 < len(lines):
            return lines[i + 1]
    return lines[0][:60] if lines else "Document"


# ---------------- PDF (text + figures) ----------------
def extract_pdf(path: Path) -> tuple[list[Segment], list[tuple[str, bytes, str, str]], int]:
    """Returns (text segments, figures[(location, bytes, media_type, page_text)], page_count)."""
    import pymupdf

    doc = pymupdf.open(path)
    segments, figures = [], []
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text")
        segments.append(Segment(location=f"p. {i}", text=text, topic=page_topic(text)))
        for img in page.get_images(full=True):
            if len(figures) >= settings().max_figures_per_doc:
                break
            try:
                info = doc.extract_image(img[0])
            except Exception:
                continue
            if info.get("width", 0) < 120 or info.get("height", 0) < 120:
                continue  # icons, bullets
            figures.append((f"Figure p. {i}", info["image"], f"image/{info.get('ext', 'png').replace('jpg', 'jpeg')}", text[:1500]))
        # Vector diagrams (drawn shapes, no embedded image): render the page if it has many drawings.
        if len(figures) < settings().max_figures_per_doc and len(page.get_drawings()) > 25:
            pix = page.get_pixmap(dpi=110)
            figures.append((f"Diagram p. {i}", pix.tobytes("png"), "image/png", text[:1500]))
    return segments, figures, len(doc)


async def caption_figures(figures: list[tuple[str, bytes, str, str]]) -> list[Segment]:
    """Multimodal step: a vision model explains each diagram so it becomes searchable and citable."""
    if not figures or not enabled("vision"):
        return []
    out = []
    for location, data, media_type, context in figures:
        try:
            res = await complete(
                "vision",
                "You describe figures from course material for students. Be factual and specific; do not invent labels.",
                [{"role": "user", "content": [image_block(data, media_type), {"type": "text", "text": f"Page text for context:\n{context}\n\nExplain what this figure shows in 2-4 sentences: components, labels, and how they relate."}]}],
                max_tokens=600,
                effort="low",
            )
        except LLMUnavailable as e:
            log.warning("figure captioning stopped: %s", e)
            break
        if res.text.strip():
            out.append(Segment(location=location, text=res.text.strip(), topic=page_topic(context) + ": Figure", modality="figure"))
    return out


# ---------------- Slides ----------------
def extract_pptx(path: Path) -> tuple[list[Segment], int]:
    from pptx import Presentation

    prs = Presentation(str(path))
    segs = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            texts.append(slide.notes_slide.notes_text_frame.text)
        body = "\n".join(texts)
        if body.strip():
            segs.append(Segment(location=f"Slide {i}", text=body, topic=texts[0].split("\n")[0][:60] if texts else f"Slide {i}"))
    return segs, len(prs.slides)


# ---------------- DOCX ----------------
def extract_docx(path: Path) -> list[Segment]:
    import zipfile

    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf8", "ignore")
    paras = [re.sub(r"<[^>]+>", "", p) for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)]
    return _group_text("\n".join(p for p in paras if p.strip()), "Section")


# ---------------- Video / audio ----------------
def _ts(sec: float) -> str:
    s = int(sec)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


def _window(snippets: list[tuple[float, str]], seconds: int = 60) -> list[Segment]:
    segs, cur, start = [], [], None
    for t, text in snippets:
        if start is None:
            start = t
        cur.append(text)
        if t - start >= seconds:
            segs.append(Segment(location=_ts(start), text=" ".join(cur), modality="transcript"))
            cur, start = [], None
    if cur:
        segs.append(Segment(location=_ts(start or 0), text=" ".join(cur), modality="transcript"))
    return segs


def transcribe(path: Path) -> tuple[list[Segment], str]:
    try:
        from faster_whisper import WhisperModel
    except ImportError as e:
        raise RuntimeError("Install faster-whisper (pip install faster-whisper) to transcribe video/audio") from e
    model = WhisperModel(settings().whisper_model, compute_type="int8")
    parts, info = model.transcribe(str(path), vad_filter=True)
    snippets = [(p.start, p.text.strip()) for p in parts]
    return _window(snippets), _ts(info.duration)


def youtube_transcript(url: str) -> tuple[list[Segment], str]:
    from youtube_transcript_api import YouTubeTranscriptApi

    m = re.search(r"(?:v=|youtu\.be/|embed/|shorts/)([\w-]{6,})", url)
    if not m:
        raise ValueError("Not a YouTube URL")
    fetched = YouTubeTranscriptApi().fetch(m.group(1), languages=("en", "hi"))
    snippets = [(s.start, s.text) for s in fetched]
    end = snippets[-1][0] if snippets else 0
    return _window(snippets), _ts(end)


# ---------------- Web / text ----------------
def extract_web(url: str) -> tuple[list[Segment], str]:
    import httpx
    from bs4 import BeautifulSoup

    html = httpx.get(url, follow_redirects=True, timeout=20, headers={"User-Agent": "StudyNest/1.0"}).text
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header", "aside"]):
        t.decompose()
    title = (soup.title.string or url).strip() if soup.title else url
    segs, cur_head, cur = [], title, []
    for el in soup.find_all(["h1", "h2", "h3", "p", "li"]):
        txt = el.get_text(" ", strip=True)
        if not txt:
            continue
        if el.name in {"h1", "h2", "h3"}:
            if cur:
                segs.append(Segment(location=f"§ {cur_head[:40]}", text="\n".join(cur), topic=cur_head))
            cur_head, cur = txt, []
        else:
            cur.append(txt)
    if cur:
        segs.append(Segment(location=f"§ {cur_head[:40]}", text="\n".join(cur), topic=cur_head))
    return segs, title[:80]


def _group_text(text: str, label: str) -> list[Segment]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paras) <= 1:
        paras = [p for p in re.split(r"(?<=[.!?])\s+(?=[A-Z])", text) if p.strip()]
    segs = []
    for i in range(0, len(paras), 2):
        segs.append(Segment(location=f"{label} {i // 2 + 1}", text=" ".join(paras[i : i + 2])))
    return segs


def extract_text(text: str) -> list[Segment]:
    return _group_text(text, "Para")


async def extract_image(path: Path, media_type: str) -> list[Segment]:
    return await caption_figures([("Image", path.read_bytes(), media_type, "")])
